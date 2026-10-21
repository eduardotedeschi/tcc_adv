import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from utils.docx_variaveis import extrair_variaveis, substituir_variaveis


def _criar_docx_temporario(texto: str) -> str:
    documento = Document()
    documento.add_paragraph(texto)
    caminho = os.path.join(tempfile.gettempdir(), "teste_template.docx")
    documento.save(caminho)
    return caminho


def test_extrair_variaveis():
    caminho = _criar_docx_temporario("Prezado(a) {{nome}}, CPF {{cpf}}.")
    variaveis = extrair_variaveis(caminho)
    assert variaveis == {"nome", "cpf"}


def test_substituir_variaveis():
    caminho_entrada = _criar_docx_temporario("Prezado(a) {{nome}}, seu CPF é {{cpf}}.")
    caminho_saida = os.path.join(tempfile.gettempdir(), "teste_saida.docx")

    substituir_variaveis(caminho_entrada, caminho_saida, {"nome": "João", "cpf": "12345678900"})

    documento_resultado = Document(caminho_saida)
    texto_final = documento_resultado.paragraphs[0].text
    assert "João" in texto_final
    assert "12345678900" in texto_final
    assert "{{" not in texto_final
