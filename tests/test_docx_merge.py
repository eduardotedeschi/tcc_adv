import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from utils.docx_merge import gerar_secao_pedidos, unir_documentos


def _criar_docx_temporario(nome: str, texto: str) -> str:
    documento = Document()
    documento.add_paragraph(texto)
    caminho = os.path.join(tempfile.gettempdir(), nome)
    documento.save(caminho)
    return caminho


def test_unir_documentos():
    tese1 = _criar_docx_temporario("tese1.docx", "Texto da primeira tese.")
    tese2 = _criar_docx_temporario("tese2.docx", "Texto da segunda tese.")
    saida = os.path.join(tempfile.gettempdir(), "peticao_unida.docx")

    unir_documentos([tese1, tese2], saida)

    resultado = Document(saida)
    textos = [p.text for p in resultado.paragraphs if p.text.strip()]
    assert "Texto da primeira tese." in textos
    assert "Texto da segunda tese." in textos


def test_unir_documentos_lista_vazia_lanca_erro():
    try:
        unir_documentos([], "saida.docx")
        assert False, "deveria ter lançado ValueError"
    except ValueError:
        pass


def test_gerar_secao_pedidos():
    texto = gerar_secao_pedidos(["Tese A", "Tese B"])
    assert "1. Seja acolhido o pedido referente a: Tese A;" in texto
    assert "2. Seja acolhido o pedido referente a: Tese B;" in texto
