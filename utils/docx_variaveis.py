"""
Substituição de variáveis {{campo}} em templates .docx, usando python-docx.
Atende RF04. Convenção fixada no projeto (ver memória do time): variáveis
usam a sintaxe {{campo}}, ex.: {{nome_cliente}}, {{cpf_cliente}}.
"""
import re
from typing import Dict

from docx import Document

PADRAO_VARIAVEL = re.compile(r"\{\{\s*(\w+)\s*\}\}")


def extrair_variaveis(caminho_arquivo: str) -> set:
    """Lê um .docx e devolve o conjunto de nomes de variáveis encontradas."""
    documento = Document(caminho_arquivo)
    variaveis = set()
    for paragrafo in documento.paragraphs:
        variaveis.update(PADRAO_VARIAVEL.findall(paragrafo.text))
    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                variaveis.update(PADRAO_VARIAVEL.findall(celula.text))
    return variaveis


def substituir_variaveis(caminho_entrada: str, caminho_saida: str, valores: Dict[str, str]) -> None:
    """
    Substitui {{campo}} pelo valor correspondente em `valores` e salva o
    resultado em `caminho_saida`. Variáveis sem valor correspondente são
    mantidas como estão (para o usuário perceber o que faltou preencher).
    """
    documento = Document(caminho_entrada)

    def _substituir_em_paragrafo(paragrafo):
        texto_completo = "".join(run.text for run in paragrafo.runs)
        if "{{" not in texto_completo:
            return
        texto_substituido = PADRAO_VARIAVEL.sub(
            lambda m: str(valores.get(m.group(1), m.group(0))), texto_completo
        )
        if paragrafo.runs:
            paragrafo.runs[0].text = texto_substituido
            for run in paragrafo.runs[1:]:
                run.text = ""

    for paragrafo in documento.paragraphs:
        _substituir_em_paragrafo(paragrafo)

    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    _substituir_em_paragrafo(paragrafo)

    documento.save(caminho_saida)
